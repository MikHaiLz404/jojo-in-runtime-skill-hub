#!/usr/bin/env python3
"""
generate_xlsx.py — Comprehensive PR/Commit Report for Unity projects.

Generates an adaptive .xlsx report based on what file types actually changed:
  - Balance / Game Data  (.asset files with Fixed-Point RawValues)
  - Code / Feature       (.cs files)
  - Prefab / Scene       (.prefab, .unity files)
  - Animation / VFX      (.anim, .controller, .vfxgraph files)

Usage:
    # Auto-named output (recommended):
    python generate_xlsx.py --repo <path> --commit <hash> --output-dir <dir>

    # Explicit output path:
    python generate_xlsx.py --repo <path> --commit <hash> --output <file.xlsx>

    # Detect file types only (no xlsx):
    python generate_xlsx.py --repo <path> --commit <hash> --detect
"""

import argparse
import subprocess
import re
import os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ═══════════════════════════════════════════════════════════════════════════════
# COLOURS & STYLE HELPERS
# ═══════════════════════════════════════════════════════════════════════════════
DARK     = "1F2D3D"
BALANCE  = "922B21"   # deep red — balance
CODE     = "1A5276"   # deep blue — code
PREFAB   = "1E8449"   # deep green — prefab/scene
ANIM     = "6C3483"   # purple — animation/VFX
QA_COL   = "17202A"   # near-black — QA
ALT      = "F4F6F9"
WHITE    = "FFFFFF"
UP_BG    = "FDEDEC"   # light red = value up
DOWN_BG  = "D5F5E3"   # light green = value down
NEW_BG   = "EBF5FB"   # light blue = new field
CODE_BG  = "EAF2FF"   # code change row bg
DEL_BG   = "FDEDEC"   # deleted file
ADD_BG   = "D5F5E3"   # added file
MOD_BG   = "FEF9E7"   # modified file
PRI_BG   = {"🔴 High": "FADBD8", "🟡 Medium": "FEF9E7", "🟢 Low": "EAFAF1"}

def _fill(c):
    return PatternFill("solid", start_color=c, fgColor=c)

def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def _hcell(ws, r, c, val, bg, fc="FFFFFF", size=10):
    cell = ws.cell(r, c, val)
    cell.font = Font(name="Arial", bold=True, size=size, color=fc)
    cell.fill = _fill(bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border()

def _dcell(ws, r, c, val, bg=None, center=False, bold=False):
    cell = ws.cell(r, c, val)
    cell.font = Font(name="Arial", size=10, bold=bold)
    if bg:
        cell.fill = _fill(bg)
    cell.alignment = Alignment(horizontal="center" if center else "left",
                               vertical="center", wrap_text=True)
    cell.border = _border()

def _title(ws, r, ncols, text, bg, size=13):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    c = ws.cell(r, 1, text)
    c.font = Font(name="Arial", bold=True, size=size, color="FFFFFF")
    c.fill = _fill(bg)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r].height = 34

def _sec(ws, r, ncols, text, bg):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    c = ws.cell(r, 1, text)
    c.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    c.fill = _fill(bg)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = _border()
    ws.row_dimensions[r].height = 22


# ═══════════════════════════════════════════════════════════════════════════════
# GIT HELPERS
# ═══════════════════════════════════════════════════════════════════════════════
def _git(repo, *args):
    return subprocess.run(["git"] + list(args), cwd=repo,
                          capture_output=True, text=True).stdout


def get_commit_info(repo, hash_):
    show = _git(repo, "show", hash_, "--stat")
    info = {
        "hash": hash_[:7], "message": "", "author": "", "date": "",
        "parent1": None, "parent2": None, "is_merge": False,
        "ticket": "", "branch": "", "approved_by": "",
    }
    body_lines, in_body = [], False
    for line in show.splitlines():
        if line.startswith("Merge:"):
            parts = line.split()
            info["parent1"], info["parent2"] = parts[1], parts[2]
            info["is_merge"] = True
        elif line.startswith("Author:"):
            info["author"] = line[8:].strip()
        elif line.startswith("Date:"):
            info["date"] = line[6:].strip()
            in_body = True
        elif in_body and "|" not in line and not re.match(r"\s*\d+ files? changed", line):
            s = line.strip()
            if s:
                body_lines.append(s)
    full = "\n".join(body_lines)
    m = re.search(r"\[([A-Z]+-\d+)\]", full)
    info["ticket"] = m.group(1) if m else ""
    m2 = re.search(r"Merged in (\S+)", full)
    info["branch"] = m2.group(1) if m2 else ""
    m3 = re.search(r"Approved-by:\s*(.+)", full)
    info["approved_by"] = m3.group(1).strip() if m3 else ""
    info["message"] = body_lines[0] if body_lines else ""
    return info


def get_diff(repo, info):
    if info["is_merge"]:
        return _git(repo, "diff", info["parent1"], info["parent2"])
    return _git(repo, "show", info["hash"], "-p")


# ═══════════════════════════════════════════════════════════════════════════════
# FILE TYPE CLASSIFICATION
# ═══════════════════════════════════════════════════════════════════════════════
BALANCE_EXTS = {".asset"}
CODE_EXTS    = {".cs"}
PREFAB_EXTS  = {".prefab", ".unity"}
ANIM_EXTS    = {".anim", ".controller", ".playable"}
VFX_EXTS     = {".vfxgraph", ".shadergraph", ".shadersubgraph"}

def classify_files(diff_text):
    """Scan diff headers and group changed files by type."""
    result = defaultdict(list)  # category -> [file paths]
    for line in diff_text.splitlines():
        if line.startswith("+++ b/"):
            path = line[6:]
            ext = os.path.splitext(path)[1].lower()
            if ext in BALANCE_EXTS:
                result["balance"].append(path)
            elif ext in CODE_EXTS:
                result["code"].append(path)
            elif ext in PREFAB_EXTS:
                result["prefab"].append(path)
            elif ext in ANIM_EXTS:
                result["anim"].append(path)
            elif ext in VFX_EXTS:
                result["vfx"].append(path)
            else:
                result["other"].append(path)
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# BALANCE PARSER (.asset files — Fixed-Point RawValue)
# ═══════════════════════════════════════════════════════════════════════════════
FP = 65536.0

def _to_val(raw):
    v = raw / FP
    return str(int(v)) if v == int(v) else f"{v:.2f}".rstrip("0").rstrip(".")

def _pct(old, new):
    if old == 0: return "New"
    delta = (new - old) / abs(old) * 100
    return f"{'+' if delta > 0 else ''}{delta:.0f}%"

def _priority_balance(pct_str, field):
    if "Resistance" in field or "HitStunDamageLimit" in field:
        return "🔴 High"
    if pct_str in ("New", "Rename"):
        return "🟡 Medium"
    try:
        val = abs(float(pct_str.replace("%", "").replace("+", "")))
    except ValueError:
        return "🟡 Medium"
    return "🔴 High" if val >= 40 else "🟡 Medium" if val >= 10 else "🟢 Low"

CATEGORY_MAP = {
    "MaxHealth": "Enemy HP", "Cooldown": "Attack Cooldown",
    "Weight": "Attack Weight", "KnockBackForce": "KnockBack",
    "KnockUpForce": "KnockBack", "HitStunTime": "HitStun",
    "HitStunDamageMultiplier": "HitStun", "HitStunDamageLimit": "Resistances",
    "HitStunResistance": "Resistances", "KnockBackResistance": "Resistances",
    "MaxSpeed": "Speed / Movement", "DefaultAttackCooldown": "AI Behavior",
    "FindCloset": "AI Behavior", "MaxDuration": "Mission Settings",
    "MaxDurationPerPlayerCount": "Mission Settings",
    "MaxEnergy": "New System Fields", "BaseEnergyGain": "New System Fields",
    "EnergyGainMultiplier": "New System Fields", "IsCritical": "New System Fields",
    "ForceChargeSuccess": "New System Fields", "ChargeScalingData": "New System Fields",
    "AutomateAttackHitBoxDatas": "New System Fields",
}
NEW_FIELDS = {"AutomateAttackHitBoxDatas", "IsCritical", "ForceChargeSuccess",
              "MaxEnergy", "BaseEnergyGain", "EnergyGainMultiplier", "ChargeScalingData",
              "PhaseData", "ConditionTriggerAttacks", "DefaultTargetType"}
KNOWN_FIELDS = set(CATEGORY_MAP.keys()) | NEW_FIELDS


class BalanceChange:
    def __init__(self, asset, field, old_raw, new_raw):
        self.asset = asset
        self.short = _asset_short(asset)
        self.field = field
        self.old_val = _to_val(old_raw) if old_raw is not None else "—"
        self.new_val = _to_val(new_raw) if new_raw is not None else "—"
        self.pct = _pct(old_raw, new_raw) if (old_raw is not None and new_raw is not None) else "New"
        self.category = next((v for k, v in CATEGORY_MAP.items() if k in field), "Other")
        self.priority = _priority_balance(self.pct, field)

    def bg(self):
        if self.pct == "New": return NEW_BG
        try:
            v = float(self.pct.replace("%", "").replace("+", ""))
            return UP_BG if v > 0 else DOWN_BG if v < 0 else ALT
        except ValueError:
            return ALT


def _asset_short(path):
    name = os.path.basename(path).replace(".asset", "")
    parts = path.split("/")
    char = next((parts[i + 1] for i, p in enumerate(parts)
                 if p in ("Player", "Enemy", "Behavior") and i + 1 < len(parts)), "")
    return f"{char}/{name}" if char else name


def parse_balance(diff_text):
    changes, seen, current_file = [], set(), None
    last_field, pending_old = None, {}
    plus_fields, minus_fields = set(), set()

    def emit(f, field, old, new):
        k = (f, field)
        if k not in seen and old != new:
            seen.add(k)
            changes.append(BalanceChange(f, field, old, new))

    def flush_new():
        for field in plus_fields - minus_fields:
            if field in NEW_FIELDS and current_file:
                emit(current_file, field, None, 0)

    for line in diff_text.splitlines():
        if line.startswith("diff --git"):
            flush_new()
            current_file = last_field = None
            pending_old.clear(); plus_fields.clear(); minus_fields.clear()
            continue
        if line.startswith("+++ b/"):
            p = line[6:]
            current_file = p if p.endswith(".asset") else None
            continue
        if current_file is None or line.startswith("---") or line.startswith("+++"):
            continue
        prefix = "-" if line.startswith("-") else "+" if line.startswith("+") else " "
        body = line[1:].strip() if prefix in "+-" else line.strip()
        mf = re.match(r"^([A-Za-z]\w+):", body)
        if mf and not body.startswith("RawValue"):
            last_field = mf.group(1)
            if prefix == "-": minus_fields.add(last_field)
            elif prefix == "+": plus_fields.add(last_field)
            continue
        mr = re.match(r"^RawValue:\s*(-?\d+)", body)
        if mr and last_field and last_field in KNOWN_FIELDS:
            raw = int(mr.group(1))
            if prefix == "-":
                pending_old[last_field] = raw
            elif prefix == "+":
                emit(current_file, last_field, pending_old.pop(last_field, None), raw)

    flush_new()
    return changes


# ═══════════════════════════════════════════════════════════════════════════════
# CODE PARSER (.cs files)
# ═══════════════════════════════════════════════════════════════════════════════
class CodeChange:
    def __init__(self, path, lines_added, lines_removed, classes, change_type):
        self.path = path
        self.short = _cs_short(path)
        self.system = _cs_system(path)
        self.lines_added = lines_added
        self.lines_removed = lines_removed
        self.classes = classes
        self.type = change_type   # "Modified" | "Added" | "Deleted"
        self.priority = _priority_code(self, lines_added + lines_removed)

def _cs_short(path):
    return os.path.basename(path)

def _cs_system(path):
    """Infer system/domain from file path."""
    parts = path.replace("\\", "/").split("/")
    # Look for meaningful folder names
    skip = {"Assets", "Scripts", "RP", "src", "Source", "Packages"}
    meaningful = [p for p in parts[:-1] if p not in skip and p and not p.startswith(".")]
    return "/".join(meaningful[-2:]) if len(meaningful) >= 2 else (meaningful[-1] if meaningful else "Unknown")

def _priority_code(change, total_lines):
    if change.type in ("Added", "Deleted"): return "🟡 Medium"
    if total_lines >= 100: return "🔴 High"
    if total_lines >= 20:  return "🟡 Medium"
    return "🟢 Low"

def parse_code(diff_text):
    changes = []
    current = None
    added = removed = 0
    classes = set()

    for line in diff_text.splitlines():
        if line.startswith("diff --git"):
            if current:
                t = "Added" if removed == 0 and added > 0 else \
                    "Deleted" if added == 0 and removed > 0 else "Modified"
                changes.append(CodeChange(current, added, removed, sorted(classes), t))
            current = None; added = removed = 0; classes = set()
        elif line.startswith("+++ b/"):
            p = line[6:]
            if p.endswith(".cs"):
                current = p
        elif current:
            if line.startswith("+") and not line.startswith("+++"):
                added += 1
                m = re.search(r"\b(?:public|private|protected|internal)\s+(?:(?:static|abstract|sealed|partial)\s+)*class\s+(\w+)", line)
                if m: classes.add(m.group(1))
            elif line.startswith("-") and not line.startswith("---"):
                removed += 1

    if current:
        t = "Added" if removed == 0 and added > 0 else \
            "Deleted" if added == 0 and removed > 0 else "Modified"
        changes.append(CodeChange(current, added, removed, sorted(classes), t))

    return changes


# ═══════════════════════════════════════════════════════════════════════════════
# PREFAB / SCENE PARSER (.prefab, .unity)
# ═══════════════════════════════════════════════════════════════════════════════
class PrefabChange:
    def __init__(self, path, change_type, objects_added, objects_removed):
        self.path = path
        self.name = os.path.basename(path).replace(".prefab", "").replace(".unity", "")
        self.ext = os.path.splitext(path)[1]
        self.type = change_type
        self.objects_added = objects_added
        self.objects_removed = objects_removed
        self.priority = "🔴 High" if change_type in ("Added", "Deleted") else "🟡 Medium"

def parse_prefabs(diff_text):
    changes = []
    current = None
    obj_added = obj_removed = []
    names_plus, names_minus = set(), set()
    is_new_file = False

    for line in diff_text.splitlines():
        if line.startswith("diff --git"):
            if current:
                t = "Added" if is_new_file else "Deleted" if (not names_plus and names_minus) else "Modified"
                changes.append(PrefabChange(current, t,
                    sorted(names_plus - names_minus),
                    sorted(names_minus - names_plus)))
            current = None; names_plus = set(); names_minus = set(); is_new_file = False
        elif line.startswith("+++ b/"):
            p = line[6:]
            if any(p.endswith(ext) for ext in (".prefab", ".unity")):
                current = p
        elif line.startswith("--- /dev/null"):
            is_new_file = True
        elif current:
            # Extract GameObject names from Unity YAML
            m = re.match(r"^([+-])\s*m_Name:\s*(.+)", line)
            if m:
                name = m.group(2).strip()
                if name and name not in ("", "New Game Object"):
                    if m.group(1) == "+": names_plus.add(name)
                    else: names_minus.add(name)

    if current:
        t = "Added" if is_new_file else "Modified"
        changes.append(PrefabChange(current, t,
            sorted(names_plus - names_minus),
            sorted(names_minus - names_plus)))

    return changes


# ═══════════════════════════════════════════════════════════════════════════════
# ANIMATION / VFX PARSER (.anim, .controller, .vfxgraph, etc.)
# ═══════════════════════════════════════════════════════════════════════════════
class AnimChange:
    def __init__(self, path, change_type):
        self.path = path
        self.name = os.path.basename(path)
        ext = os.path.splitext(path)[1].lower()
        self.file_type = {
            ".anim": "Animation Clip", ".controller": "Animator Controller",
            ".playable": "Playable Graph", ".vfxgraph": "VFX Graph",
            ".shadergraph": "Shader Graph", ".shadersubgraph": "Shader Subgraph",
        }.get(ext, ext)
        self.type = change_type
        self.priority = "🟡 Medium" if change_type == "Modified" else "🔴 High"

def parse_anims(diff_text):
    changes = []
    seen_new = set()

    for line in diff_text.splitlines():
        if line.startswith("diff --git"):
            is_new = False
        elif line.startswith("--- /dev/null"):
            is_new = True
        elif line.startswith("+++ b/"):
            p = line[6:]
            ext = os.path.splitext(p)[1].lower()
            if ext in (ANIM_EXTS | VFX_EXTS):
                t = "Added" if is_new else "Modified"
                changes.append(AnimChange(p, t))

    return changes


# ═══════════════════════════════════════════════════════════════════════════════
# QA CHECKLIST — unified across all change types
# ═══════════════════════════════════════════════════════════════════════════════
def build_qa(balance, code, prefabs, anims):
    items = []   # (source, asset/system, test_description, priority)

    # Balance
    seen = set()
    for c in balance:
        k = (c.short, c.field)
        if k in seen: continue
        seen.add(k)
        pri = c.priority
        if c.category == "Enemy HP":
            items.append(("Balance", c.short, f"MaxHealth {c.old_val} → {c.new_val} ({c.pct}) — ตรวจสอบ fight duration และ balance", pri))
        elif c.category == "HitStun":
            items.append(("Balance", c.short, f"{c.field}: {c.old_val} → {c.new_val} — ตรวจสอบ combo feel และ HitStun window", pri))
        elif c.category == "KnockBack":
            items.append(("Balance", c.short, f"{c.field}: {c.old_val} → {c.new_val} ({c.pct}) — ตรวจสอบระยะ knockback", pri))
        elif c.category == "Attack Cooldown":
            items.append(("Balance", c.short, f"Cooldown {c.old_val} → {c.new_val} ({c.pct}) — ตรวจสอบ attack frequency", pri))
        elif c.category == "Attack Weight":
            items.append(("Balance", c.short, f"Weight {c.old_val} → {c.new_val} ({c.pct}) — ตรวจสอบ attack selection pattern", pri))
        elif c.category == "Resistances":
            items.append(("Balance", c.short, f"{c.field}: {c.old_val} → {c.new_val} — ตรวจสอบว่าศัตรู stun/knockback ง่ายขึ้น", "🔴 High"))
        elif c.category == "Mission Settings":
            items.append(("Balance", c.short, f"Duration {c.old_val}s → {c.new_val}s — ตรวจสอบว่า complete ได้ใน {c.new_val}s", "🔴 High"))
        elif c.category == "AI Behavior":
            items.append(("Balance", c.short, f"{c.field}: {c.old_val} → {c.new_val} ({c.pct}) — ตรวจสอบ AI behavior pattern", pri))
        elif c.category == "Speed / Movement":
            items.append(("Balance", c.short, f"MaxSpeed {c.old_val} → {c.new_val} ({c.pct}) — ตรวจสอบ movement feel", pri))
        elif c.category == "New System Fields":
            items.append(("Balance", c.short, f"[New] {c.field} = {c.new_val} — ยืนยันว่า field ใหม่ไม่กระทบ gameplay เดิม", "🟡 Medium"))

    # Code
    for c in code:
        pri = c.priority
        class_str = f"({', '.join(c.classes[:3])})" if c.classes else ""
        if c.type == "Added":
            items.append(("Code", c.system, f"[New] {c.short} {class_str} — ทดสอบ feature ใหม่ทั้งหมด", "🔴 High"))
        elif c.type == "Deleted":
            items.append(("Code", c.system, f"[Removed] {c.short} — ตรวจสอบว่าไม่มี dependency หลงเหลือ", "🔴 High"))
        else:
            items.append(("Code", c.system, f"{c.short} {class_str} (+{c.lines_added}/-{c.lines_removed} lines) — ตรวจสอบ behavior ที่เปลี่ยนไป", pri))

    # Prefab / Scene
    for p in prefabs:
        label = "Scene" if p.ext == ".unity" else "Prefab"
        if p.type == "Added":
            items.append((label, p.name, f"[New {label}] ตรวจสอบว่า appear ถูกต้องใน scene", "🔴 High"))
        else:
            obj_str = ""
            if p.objects_added:
                obj_str += f"  +{', '.join(p.objects_added[:3])}"
            if p.objects_removed:
                obj_str += f"  -{', '.join(p.objects_removed[:3])}"
            items.append((label, p.name, f"Verify {label} configuration{obj_str}".strip(), p.priority))

    # Animation / VFX
    for a in anims:
        if a.type == "Added":
            items.append((a.file_type, a.name, f"[New] Play-test ใหม่ทั้งหมด: {a.name}", "🔴 High"))
        else:
            items.append((a.file_type, a.name, f"Play-test หลังแก้ไข: {a.name}", a.priority))

    # Deduplicate
    final, seen2 = [], set()
    for item in items:
        k = item[:3]
        if k not in seen2:
            seen2.add(k)
            final.append(item)
    return final


# ═══════════════════════════════════════════════════════════════════════════════
# RISK ASSESSMENT
# ═══════════════════════════════════════════════════════════════════════════════
def assess_risk(balance, code, prefabs, anims, file_counts):
    total = sum(file_counts.values())
    high_balance = sum(1 for c in balance if c.priority == "🔴 High")
    has_new_prefab = any(p.type == "Added" for p in prefabs)
    has_new_code = any(c.type == "Added" for c in code)
    big_code = any(c.lines_added + c.lines_removed >= 100 for c in code)

    if high_balance >= 3 or has_new_prefab or big_code or total >= 30:
        return "🔴 High", "Large scope or significant value changes"
    if high_balance >= 1 or has_new_code or len(code) >= 3 or total >= 10:
        return "🟡 Medium", "Moderate changes — targeted testing required"
    return "🟢 Low", "Small focused change — smoke test sufficient"


# ═══════════════════════════════════════════════════════════════════════════════
# XLSX BUILDER
# ═══════════════════════════════════════════════════════════════════════════════
def build_xlsx(info, balance, code, prefabs, anims, output):
    wb = Workbook()
    file_counts = {k: len(v) for k, v in {
        "balance": balance, "code": code, "prefab": prefabs, "anim": anims
    }.items() if v}

    # ── Sheet 1: PR Overview ───────────────────────────────────────────────────
    ws = wb.active
    ws.title = "📋 PR Overview"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 54

    ticket = info.get("ticket", "")
    title_text = f"📋 {ticket} — PR Report" if ticket else "📋 Commit Report"
    _title(ws, 1, 2, title_text, DARK)

    meta = [
        ("Commit",      info["hash"]),
        ("Branch",      info.get("branch", "")),
        ("Date",        info["date"]),
        ("Author",      info["author"]),
        ("Approved by", info.get("approved_by", "")),
        ("Ticket",      ticket),
    ]
    for i, (k, v) in enumerate(meta, 2):
        ws.row_dimensions[i].height = 20
        ck = ws.cell(i, 1, k); ck.font = Font(name="Arial", bold=True, size=10)
        ck.fill = _fill(ALT); ck.border = _border()
        cv = ws.cell(i, 2, v); cv.font = Font(name="Arial", size=10); cv.border = _border()

    risk_level, risk_note = assess_risk(balance, code, prefabs, anims, file_counts)
    ws.row_dimensions[9].height = 8
    _hcell(ws, 10, 1, "Risk", DARK); _hcell(ws, 10, 2, "Note", DARK)
    ws.row_dimensions[10].height = 20
    _dcell(ws, 11, 1, risk_level, MOD_BG, center=True)
    _dcell(ws, 11, 2, risk_note, MOD_BG)

    ws.row_dimensions[13].height = 8
    _hcell(ws, 14, 1, "Change Type", DARK); _hcell(ws, 14, 2, "Summary", DARK)
    row = 15
    type_info = [
        (balance, "⚖️ Balance / Game Data", f"{len(balance)} value changes in .asset files"),
        (code,    "💻 Code / Feature",       f"{len(code)} .cs files (+{sum(c.lines_added for c in code)}/-{sum(c.lines_removed for c in code)} lines)"),
        (prefabs, "🎮 Prefab / Scene",        f"{len(prefabs)} prefab/scene files changed"),
        (anims,   "🎬 Animation / VFX",       f"{len(anims)} animation/VFX files changed"),
    ]
    for items, label, summary in type_info:
        if items:
            ws.row_dimensions[row].height = 22
            _dcell(ws, row, 1, label, ALT if row % 2 == 0 else WHITE, bold=True)
            _dcell(ws, row, 2, summary, ALT if row % 2 == 0 else WHITE)
            row += 1

    # ── Sheet 2: Balance Changes ───────────────────────────────────────────────
    if balance:
        ws2 = wb.create_sheet("⚖️ Balance Changes")
        ws2.sheet_view.showGridLines = False
        for col, w in zip("ABCDEF", [32, 26, 12, 12, 12, 36]):
            ws2.column_dimensions[col].width = w
        _title(ws2, 1, 6, f"⚖️ Balance / Game Data Changes — {ticket}", BALANCE)
        for ci, v in enumerate(["Asset", "Field", "Before", "After", "% Change", "Category"], 1):
            _hcell(ws2, 3, ci, v, BALANCE)
        for i, c in enumerate(balance, 4):
            bg = c.bg()
            ws2.row_dimensions[i].height = 20
            _dcell(ws2, i, 1, c.short, ALT if i % 2 == 0 else WHITE)
            _dcell(ws2, i, 2, c.field, bg)
            _dcell(ws2, i, 3, c.old_val, bg, center=True)
            _dcell(ws2, i, 4, c.new_val, bg, center=True)
            _dcell(ws2, i, 5, c.pct, bg, center=True)
            _dcell(ws2, i, 6, c.category, bg)

    # ── Sheet 3: Code Changes ──────────────────────────────────────────────────
    if code:
        ws3 = wb.create_sheet("💻 Code Changes")
        ws3.sheet_view.showGridLines = False
        for col, w in zip("ABCDEF", [20, 30, 10, 10, 10, 36]):
            ws3.column_dimensions[col].width = w
        _title(ws3, 1, 6, f"💻 Code / Feature Changes — {ticket}", CODE)
        for ci, v in enumerate(["System / Domain", "File", "Type", "+ Lines", "- Lines", "Classes Affected"], 1):
            _hcell(ws3, 3, ci, v, CODE)
        for i, c in enumerate(code, 4):
            bg = ADD_BG if c.type == "Added" else DEL_BG if c.type == "Deleted" else (CODE_BG if i % 2 == 0 else WHITE)
            ws3.row_dimensions[i].height = 20
            _dcell(ws3, i, 1, c.system, bg)
            _dcell(ws3, i, 2, c.short, bg)
            _dcell(ws3, i, 3, c.type, bg, center=True)
            _dcell(ws3, i, 4, c.lines_added, bg, center=True)
            _dcell(ws3, i, 5, c.lines_removed, bg, center=True)
            _dcell(ws3, i, 6, ", ".join(c.classes[:5]) if c.classes else "—", bg)

    # ── Sheet 4: Prefab / Scene Changes ───────────────────────────────────────
    if prefabs:
        ws4 = wb.create_sheet("🎮 Prefab+Scene")
        ws4.sheet_view.showGridLines = False
        for col, w in zip("ABCDE", [28, 12, 28, 28, 14]):
            ws4.column_dimensions[col].width = w
        _title(ws4, 1, 5, f"🎮 Prefab / Scene Changes — {ticket}", PREFAB)
        for ci, v in enumerate(["Name", "Type", "Objects Added", "Objects Removed", "Priority"], 1):
            _hcell(ws4, 3, ci, v, PREFAB)
        for i, p in enumerate(prefabs, 4):
            bg = ADD_BG if p.type == "Added" else DEL_BG if p.type == "Deleted" else (ALT if i % 2 == 0 else WHITE)
            ws4.row_dimensions[i].height = 22
            _dcell(ws4, i, 1, p.name, bg)
            _dcell(ws4, i, 2, p.type, bg, center=True)
            _dcell(ws4, i, 3, ", ".join(p.objects_added[:5]) or "—", bg)
            _dcell(ws4, i, 4, ", ".join(p.objects_removed[:5]) or "—", bg)
            _dcell(ws4, i, 5, p.priority, bg, center=True)

    # ── Sheet 5: Animation / VFX Changes ──────────────────────────────────────
    if anims:
        ws5 = wb.create_sheet("🎬 Animation+VFX")
        ws5.sheet_view.showGridLines = False
        for col, w in zip("ABCD", [40, 22, 12, 14]):
            ws5.column_dimensions[col].width = w
        _title(ws5, 1, 4, f"🎬 Animation / VFX Changes — {ticket}", ANIM)
        for ci, v in enumerate(["File Name", "Type", "Change", "Priority"], 1):
            _hcell(ws5, 3, ci, v, ANIM)
        for i, a in enumerate(anims, 4):
            bg = ADD_BG if a.type == "Added" else (ALT if i % 2 == 0 else WHITE)
            ws5.row_dimensions[i].height = 20
            _dcell(ws5, i, 1, a.name, bg)
            _dcell(ws5, i, 2, a.file_type, bg)
            _dcell(ws5, i, 3, a.type, bg, center=True)
            _dcell(ws5, i, 4, a.priority, bg, center=True)

    # ── Sheet 6: QA Checklist ─────────────────────────────────────────────────
    ws6 = wb.create_sheet("✅ QA Checklist")
    ws6.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", [5, 22, 16, 44, 16]):
        ws6.column_dimensions[col].width = w
    _title(ws6, 1, 5, f"✅ QA Checklist — {ticket}", QA_COL)
    for ci, v in enumerate(["#", "Asset / System", "Category", "Test Case", "Priority"], 1):
        _hcell(ws6, 3, ci, v, QA_COL)

    qa = build_qa(balance, code, prefabs, anims)
    for i, (source, asset, test, pri) in enumerate(qa, 4):
        bg = PRI_BG.get(pri, WHITE)
        ws6.row_dimensions[i].height = 26
        _dcell(ws6, i, 1, str(i - 3), bg, center=True)
        _dcell(ws6, i, 2, asset, bg)
        _dcell(ws6, i, 3, source, bg)
        _dcell(ws6, i, 4, test, bg)
        _dcell(ws6, i, 5, pri, bg, center=True)

    # Legend
    lr = len(qa) + 6
    ws6.merge_cells(f"A{lr}:E{lr}")
    lc = ws6.cell(lr, 1, "สี: 🔴 High  |  🟡 Medium  |  🟢 Low")
    lc.font = Font(name="Arial", bold=True, size=9)
    lc.fill = _fill(ALT)
    lc.alignment = Alignment(horizontal="left", vertical="center")

    wb.save(output)
    return output


# ═══════════════════════════════════════════════════════════════════════════════
# AUTO FILENAME (naming convention)
# ═══════════════════════════════════════════════════════════════════════════════
def auto_filename(info, output_dir):
    raw_date = info.get("date", "")
    date_str = "00000000"
    month_map = {"Jan":"01","Feb":"02","Mar":"03","Apr":"04","May":"05","Jun":"06",
                 "Jul":"07","Aug":"08","Sep":"09","Oct":"10","Nov":"11","Dec":"12"}
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", raw_date)
    if m:
        date_str = f"{m.group(1)}{m.group(2)}{m.group(3)}"
    else:
        m2 = re.search(r"\w{3}\s+(\w{3})\s+(\d{1,2})\s+[\d:]+\s+(\d{4})", raw_date)
        if m2:
            date_str = f"{m2.group(3)}{month_map.get(m2.group(1),'00')}{m2.group(2).zfill(2)}"
    ticket = info.get("ticket", "").replace("/", "-")
    commit_short = info.get("hash", "unknown")[:7]
    filename = f"{ticket}_{date_str}.xlsx" if ticket else f"commit_{commit_short}_{date_str}.xlsx"
    return os.path.join(output_dir, filename)


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════
def main():
    parser = argparse.ArgumentParser(
        description="Generate a comprehensive PR/commit report .xlsx for Unity projects."
    )
    parser.add_argument("--repo",   required=True, help="Path to git repository")
    parser.add_argument("--commit", required=True, help="Commit hash to analyse")
    parser.add_argument("--detect", action="store_true",
                        help="Only detect file types changed (no xlsx output)")

    out_group = parser.add_mutually_exclusive_group()
    out_group.add_argument("--output",     help="Full output file path")
    out_group.add_argument("--output-dir", help="Output directory (auto-named file)")
    args = parser.parse_args()

    print(f"Reading commit {args.commit}...")
    info = get_commit_info(args.repo, args.commit)
    print(f"  Ticket : {info['ticket']}")
    print(f"  Author : {info['author']}")
    print(f"  Date   : {info['date']}")
    print(f"  Merge  : {info['is_merge']}")

    print("Getting diff...")
    diff = get_diff(args.repo, info)
    print(f"  Diff size: {len(diff):,} chars")

    # Classify files
    file_map = classify_files(diff)
    print("\nFILE TYPES CHANGED:")
    type_labels = {
        "balance": ".asset   → Balance / Game Data",
        "code":    ".cs      → Code / Feature",
        "prefab":  ".prefab  → Prefab / Scene",
        "anim":    ".anim    → Animation / VFX",
        "vfx":     ".vfx     → VFX / Shader",
        "other":   "(other)",
    }
    for k, label in type_labels.items():
        n = len(file_map.get(k, []))
        if n > 0:
            print(f"  {label}: {n} files")

    if args.detect:
        return

    if not args.output and not args.output_dir:
        parser.error("Provide --output or --output-dir (or use --detect for inspection only)")

    # Parse each type
    print("\nParsing changes...")
    balance_changes = parse_balance(diff)
    code_changes    = parse_code(diff)
    prefab_changes  = parse_prefabs(diff)
    anim_changes    = parse_anims(diff)

    print(f"  Balance : {len(balance_changes)} value changes")
    print(f"  Code    : {len(code_changes)} files")
    print(f"  Prefab  : {len(prefab_changes)} files")
    print(f"  Anim    : {len(anim_changes)} files")

    # Resolve output path
    if args.output:
        output_path = args.output
    else:
        os.makedirs(args.output_dir, exist_ok=True)
        output_path = auto_filename(info, args.output_dir)

    print(f"\nWriting {output_path}...")
    build_xlsx(info, balance_changes, code_changes, prefab_changes, anim_changes, output_path)
    print(f"OUTPUT_FILE={output_path}")
    print("Done!")


if __name__ == "__main__":
    main()
